"""Ingest the raw/source-of-record datasets into PostgreSQL.

    raw files  ->  Postgres (this script)  ->  static export  ->  SPA

Run after `make db-up` (schema applied on first init) — or pass --schema to
(re)apply db/schema.sql first. Idempotent: every table is truncated before
load, so re-running fully refreshes the database.

Sources (all committed to the repo today):
  * regions      <- webapp/data/rgint_index.json + webapp/public/data/rgint_meta.json
  * municipios   <- ILUC_NIPE/02_Spatial_Lookups/...xlsx
  * lulc_timeseries <- webapp/data/rgint/{id}.json   (15-class diagonal series)
  * transitions  <- webapp/data/rgint_matrix/{id}.json  (full 15x15 matrices)
  * pam          <- ILUC_NIPE/05_Agro_Subdivisions/PAM_RGINT_COMPLETO.csv
  * lapig_vigor  <- ILUC_NIPE/03_Pasture_Vigor_LAPIG/*.csv
  * terraclass   <- ILUC_NIPE/04_Vegetation_TerraClass/TC_*.csv
  * conab_graos  <- ILUC_NIPE/05_Agro_Subdivisions/CONAB_GRAOS_*.csv

Dropping the full 15x15 matrix JSONs (all 133 regions) into
webapp/data/rgint_matrix/ and re-running populates `transitions` for every
region — no code change needed.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

import openpyxl
from psycopg2.extras import execute_values

from common import (
    INDEX_JSON,
    MUNICIPIOS_XLSX,
    OUT,
    PAM_CSV,
    RGINT_MATRIX_DIR,
    RGINT_SERIES_DIR,
    ROOT,
)
from db import connect

GOLDEN_RGINTS = {1201, 5101}
LAPIG_DIR = ROOT / "ILUC_NIPE" / "03_Pasture_Vigor_LAPIG"
TERRACLASS_DIR = ROOT / "ILUC_NIPE" / "04_Vegetation_TerraClass"
CONAB_FILES = [
    ROOT / "ILUC_NIPE" / "05_Agro_Subdivisions" / "CONAB_GRAOS_UF_2008_2024.csv",
    ROOT / "ILUC_NIPE" / "05_Agro_Subdivisions" / "CONAB_GRAOS_CANA_UF_2008_2024.csv",
]
TC_KEY_COLS = {"CD_MUN", "SIGLA_UF", "ANO"}


def _num(value) -> float | None:
    try:
        if value in (None, "", "NaN", "nan"):
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _int(value) -> int | None:
    f = _num(value)
    return int(f) if f is not None else None


# --------------------------------------------------------------------------- #
# Per-table loaders. Each returns the rows it inserted (for logging).
# --------------------------------------------------------------------------- #
def load_regions(cur) -> int:
    index = json.loads(Path(INDEX_JSON).read_text(encoding="utf-8"))
    meta_path = OUT / "rgint_meta.json"
    meta = {}
    if meta_path.exists():
        for m in json.loads(meta_path.read_text(encoding="utf-8")):
            meta[str(m["id"])] = m
    rows = []
    for item in index:
        rid = int(item["rgint"])
        m = meta.get(item["rgint"], {})
        rows.append((
            rid, item["nome"], item["uf"], item["biome"],
            _num(m.get("area_ha")), _int(m.get("n_municipios")),
            _num(m.get("lon_centroide")), _num(m.get("lat_centroide")),
            rid in GOLDEN_RGINTS,
        ))
    execute_values(cur, (
        "INSERT INTO regions (rgint_id, nome, uf, bioma, area_ha, n_municipios, "
        "centroid_lon, centroid_lat, is_golden) VALUES %s"
    ), rows)
    return len(rows)


def load_municipios(cur) -> int:
    wb = openpyxl.load_workbook(MUNICIPIOS_XLSX, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [str(c) if c is not None else "" for c in next(ws.iter_rows(values_only=True))]
    c_mun = header.index("nome_mun")
    c_cod = header.index("CD_GEOCODI")
    c_rg = header.index("cod_rgint")
    rows = []
    seen: set[int] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        cd = _int(row[c_cod])
        rid = _int(row[c_rg])
        if cd is None or rid is None or cd in seen:
            continue
        seen.add(cd)
        rows.append((cd, row[c_mun], None, rid))
    execute_values(cur,
        "INSERT INTO municipios (cd_mun, nome_mun, uf, rgint_id) VALUES %s", rows)
    return len(rows)


def load_timeseries(cur) -> int:
    rows = []
    for path in sorted(Path(RGINT_SERIES_DIR).glob("*.json")):
        rid = int(path.stem)
        data = json.loads(path.read_text(encoding="utf-8").replace("NaN", "null"))
        for classe, by_year in data.items():
            for year, area in by_year.items():
                rows.append((rid, int(year), classe, _num(area)))
    execute_values(cur,
        "INSERT INTO lulc_timeseries (rgint_id, ano, classe, area_ha) VALUES %s",
        rows, page_size=2000)
    return len(rows)


def load_transitions(cur) -> int:
    rows = []
    if Path(RGINT_MATRIX_DIR).exists():
        for path in sorted(Path(RGINT_MATRIX_DIR).glob("*.json")):
            rid = int(path.stem)
            data = json.loads(path.read_text(encoding="utf-8").replace("NaN", "null"))
            for year, by_src in data.get("matrices", {}).items():
                for origem, dests in by_src.items():
                    for destino, area in dests.items():
                        if area is None:
                            continue
                        rows.append((rid, str(year), origem, destino, _num(area)))
    if rows:
        execute_values(cur,
            "INSERT INTO transitions (rgint_id, ano_par, origem_id, destino_id, area_ha) "
            "VALUES %s", rows, page_size=2000)
    return len(rows)


def load_pam(cur) -> int:
    rows = []
    agg: dict[tuple[int, int, str], float] = {}
    with open(PAM_CSV, encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            key = (int(r["CD_RGINT"]), int(r["ano"]), r["cultura"])
            agg[key] = agg.get(key, 0.0) + (_num(r["area_ha"]) or 0.0)
    for (rid, ano, cultura), area in agg.items():
        rows.append((rid, ano, cultura, area))
    execute_values(cur,
        "INSERT INTO pam (rgint_id, ano, cultura, area_ha) VALUES %s",
        rows, page_size=5000)
    return len(rows)


def load_lapig(cur) -> int:
    rows = []
    for path in sorted(LAPIG_DIR.glob("*.csv")):
        with open(path, encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                rows.append((
                    _int(r.get("geocod_mun")), _int(r.get("ano")),
                    r.get("bioma"), r.get("classe"), _num(r.get("area_past_ha")),
                ))
    if rows:
        execute_values(cur,
            "INSERT INTO lapig_vigor (cd_mun, ano, bioma, classe_vigor, area_past_ha) "
            "VALUES %s", rows, page_size=5000)
    return len(rows)


def load_terraclass(cur) -> int:
    rows = []
    for path in sorted(TERRACLASS_DIR.glob("TC_*.csv")):
        fonte = "CER" if "_CER_" in path.name else "AMZ"
        with open(path, encoding="utf-8") as fh:
            reader = csv.DictReader(fh)
            value_cols = [c for c in (reader.fieldnames or []) if c not in TC_KEY_COLS]
            for r in reader:
                cd = _int(r.get("CD_MUN"))
                ano = _int(r.get("ANO"))
                if cd is None or ano is None:
                    continue
                for col in value_cols:
                    area = _num(r.get(col))
                    if area is not None:
                        rows.append((cd, ano, fonte, col, area))
    if rows:
        execute_values(cur,
            "INSERT INTO terraclass (cd_mun, ano, fonte, categoria, area_ha) VALUES %s",
            rows, page_size=5000)
    return len(rows)


def load_conab(cur) -> int:
    rows = []
    for path in CONAB_FILES:
        if not path.exists():
            continue
        with open(path, encoding="utf-8-sig") as fh:
            for r in csv.DictReader(fh):
                rows.append((r.get("uf"), _int(r.get("ano")), r.get("cultura"),
                             _num(r.get("conab_ha"))))
    if rows:
        execute_values(cur,
            "INSERT INTO conab_graos (uf, ano, cultura, conab_ha) VALUES %s",
            rows, page_size=5000)
    return len(rows)


LOADERS = [
    ("regions", load_regions),
    ("municipios", load_municipios),
    ("lulc_timeseries", load_timeseries),
    ("transitions", load_transitions),
    ("pam", load_pam),
    ("lapig_vigor", load_lapig),
    ("terraclass", load_terraclass),
    ("conab_graos", load_conab),
]
# Truncate order respects FK dependencies (children first).
TRUNCATE_ORDER = [
    "indicators", "transitions", "lulc_timeseries", "pam", "municipios",
    "conab_graos", "terraclass", "lapig_vigor", "regions",
]


def main() -> None:
    ap = argparse.ArgumentParser(description="Ingest raw datasets into Postgres.")
    ap.add_argument("--schema", action="store_true",
                    help="Apply db/schema.sql before loading.")
    args = ap.parse_args()

    with connect() as conn:
        cur = conn.cursor()
        if args.schema:
            cur.execute((ROOT / "db" / "schema.sql").read_text(encoding="utf-8"))
        cur.execute("TRUNCATE " + ", ".join(TRUNCATE_ORDER) + " RESTART IDENTITY CASCADE")
        for name, loader in LOADERS:
            n = loader(cur)
            print(f"  {name:18s} {n:>8,} rows")
    print("ingest complete.")


if __name__ == "__main__":
    sys.exit(main())
