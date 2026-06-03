"""Ingest the HARVEX anchor-point deliverable into Postgres.

    HARVEX workbook + consolidado.csv  ->  harvex_transitions + harvex_trajectories

The HARVEX format (Joel Risso / HARVEX) is the GOLDEN source that will replace
the interim annual matrices. Two artifacts feed two tables:

  * ``Lookup`` sheet  (REG_ID, periodo, orig, dest, ha)
        -> harvex_transitions  (2-step flows per GTAP period)
  * ``consolidado.csv`` (REG_ID, codigo_trajetoria, count_pixels)
        -> harvex_trajectories (3-step pixel paths 2008->2017->2024)

Class codes (1..15) map positionally onto the canonical CLASS_ORDER labels —
HARVEX uses the same 15-class system as the rest of the stack. The ``periodo``
column arrives as an int (``0817`` -> ``817``); we zero-pad to 4 digits before
mapping, otherwise the period join silently breaks.

Run AFTER scripts/ingest.py (regions must exist for the FK). Idempotent: both
HARVEX tables are truncated before load.

Usage:
    python ingest_harvex.py [--xlsx PATH] [--csv PATH] [--schema]
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from psycopg2.extras import execute_values

from common import CLASS_ORDER, ROOT
from db import connect

# Placeholder sample delivery (real Joel data will land elsewhere — override via CLI).
_HARVEX_SAMPLE_DIR = Path(r"C:\Users\Lucas\Documents\ILUC_NIPE")
DEFAULT_XLSX = _HARVEX_SAMPLE_DIR / "consolidador_transicoes_lulc_ABIOVE (1).xlsx"
DEFAULT_CSV = _HARVEX_SAMPLE_DIR / "consolidado (1).csv"

PIXEL_HA = 0.09  # 30 m raster -> 0.09 ha/pixel (HARVEX '1_Parametros')

# HARVEX period code (zero-padded to 4) -> canonical GTAP period key.
PERIOD_MAP = {"0817": "2008_2017", "1724": "2017_2024", "0824": "2008_2024"}


def class_label(code) -> str:
    """Map a HARVEX class code (1..15) to its canonical CLASS_ORDER label."""
    n = int(code)
    if not 1 <= n <= 15:
        raise ValueError(f"class code out of range 1..15: {code!r}")
    return CLASS_ORDER[n - 1]


def canonical_period(raw) -> str | None:
    """'817'/817 -> '2008_2017'. Returns None for unrecognized codes."""
    return PERIOD_MAP.get(str(int(raw)).zfill(4))


def decode_trajectory(code) -> tuple[int, int, int]:
    """AABBCC trajectory code -> (class_2008, class_2017, class_2024)."""
    s = str(int(code)).zfill(6)
    return int(s[0:2]), int(s[2:4]), int(s[4:6])


def load_harvex_transitions(cur, xlsx: Path) -> int:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if "Lookup" not in wb.sheetnames:
        raise SystemExit(f"'Lookup' sheet not found in {xlsx.name}")
    ws = wb["Lookup"]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows_iter)]
    idx = {name: header.index(name) for name in ("REG_ID", "periodo", "orig", "dest", "ha")}

    # Aggregate defensively (PK is rgint+periodo+origem+destino).
    agg: dict[tuple, float] = defaultdict(float)
    skipped_period = 0
    for r in rows_iter:
        if r[idx["REG_ID"]] is None:
            continue
        periodo = canonical_period(r[idx["periodo"]])
        if periodo is None:
            skipped_period += 1
            continue
        rid = int(r[idx["REG_ID"]])
        origem = class_label(r[idx["orig"]])
        destino = class_label(r[idx["dest"]])
        ha = r[idx["ha"]]
        agg[(rid, periodo, origem, destino)] += float(ha or 0.0)
    wb.close()

    rows = [(rid, per, o, d, round(ha, 6)) for (rid, per, o, d), ha in agg.items()]
    execute_values(cur,
        "INSERT INTO harvex_transitions (rgint_id, periodo, origem_id, destino_id, area_ha) "
        "VALUES %s", rows, page_size=5000)
    if skipped_period:
        print(f"    (skipped {skipped_period} rows with unrecognized periodo)")
    return len(rows)


def load_harvex_trajectories(cur, csv_path: Path) -> int:
    agg: dict[tuple, float] = defaultdict(float)
    with open(csv_path, encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            rid = int(r["REG_ID"])
            c08, c17, c24 = decode_trajectory(r["codigo_trajetoria"])
            ha = int(r["count_pixels"]) * PIXEL_HA
            key = (rid, class_label(c08), class_label(c17), class_label(c24))
            agg[key] += ha
    rows = [(rid, a, b, c, round(ha, 6)) for (rid, a, b, c), ha in agg.items()]
    execute_values(cur,
        "INSERT INTO harvex_trajectories "
        "(rgint_id, classe_2008, classe_2017, classe_2024, area_ha) VALUES %s",
        rows, page_size=5000)
    return len(rows)


def main() -> None:
    ap = argparse.ArgumentParser(description="Ingest HARVEX anchor points into Postgres.")
    ap.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="HARVEX consolidador workbook")
    ap.add_argument("--csv", default=str(DEFAULT_CSV), help="HARVEX consolidado.csv (trajectories)")
    ap.add_argument("--schema", action="store_true",
                    help="Apply db/schema.sql before loading (rebuilds ALL tables).")
    args = ap.parse_args()

    xlsx, csv_path = Path(args.xlsx), Path(args.csv)
    for p in (xlsx, csv_path):
        if not p.exists():
            sys.exit(f"source not found: {p}")

    with connect() as conn:
        cur = conn.cursor()
        if args.schema:
            cur.execute((ROOT / "db" / "schema.sql").read_text(encoding="utf-8"))
        cur.execute("TRUNCATE harvex_trajectories, harvex_transitions RESTART IDENTITY")
        n_tr = load_harvex_transitions(cur, xlsx)
        print(f"  harvex_transitions   {n_tr:>8,} rows")
        n_tj = load_harvex_trajectories(cur, csv_path)
        print(f"  harvex_trajectories  {n_tj:>8,} rows")

        # Structural sanity (values are placeholder — we only assert shape).
        cur.execute("SELECT COUNT(DISTINCT rgint_id), COUNT(DISTINCT periodo) FROM harvex_transitions")
        n_reg, n_per = cur.fetchone()
        cur.execute("SELECT COUNT(DISTINCT rgint_id) FROM harvex_trajectories")
        (n_reg_tj,) = cur.fetchone()
        print(f"\n  transitions: {n_reg} regions x {n_per} periods (expect 133 x 3)")
        print(f"  trajectories: {n_reg_tj} regions (expect 133)")
    print("HARVEX ingest complete.")


if __name__ == "__main__":
    sys.exit(main())
