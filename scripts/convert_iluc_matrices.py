"""Build per-region 15x15 annual transition JSONs from the ABIOVE final
15-class matrices (real annual data).

    ABIOVE xlsx  ->  webapp/data/rgint_matrix/{id}.json  ->  ingest -> SPA

Source layout (``07_MATRIZES_15_CLASSES_FINAL``):
  ALL_RGINTS/ILUC_15Classes_RGINT{id}_{name}.xlsx   -- 133 regions (raw)
  GOLDEN_ILUC_15_Classes_RGINT_{id}.xlsx            -- refined 1201, 5101

Each workbook has 16 sheets (``2008_2009`` .. ``2023_2024``). Each sheet is a
16x16 grid: row 0 / col 0 hold the 15 class labels, cells [1..15][1..15] are the
``area_ha`` flowing origem(row) -> destino(col) over that year pair. The col-15
header is corrupted in the ALL_RGINTS files, so columns are mapped positionally
(col j -> class j); row labels are reliable and matched against ``CLASS_ORDER``.

Output JSON matches the shape the ingest/SPA already expect (same as the
committed 5101.json): the matrix key is the END year of each sheet
(``2008_2009`` -> ``"2009"``), and ``"2008"`` is the base-year stock as a
diagonal (row-sums of the 2008_2009 matrix = total 2008 area per class).

Usage:
    python convert_iluc_matrices.py [SRC_DIR]
    # or set ILUC_MATRIX_SRC_DIR; defaults to the local ABIOVE_SOJA_2026 path.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path

import openpyxl

from common import CLASS_ORDER, INDEX_JSON, RGINT_MATRIX_DIR

DEFAULT_SRC = Path(
    r"C:\Users\Lucas\Documents\ABIOVE_SOJA_2026"
    r"\05_FINAL_INTEGRATION_DATA\07_MATRIZES_15_CLASSES_FINAL"
)
GOLDEN_IDS = {"1201", "5101"}

# Map a class label's leading number ("7 - Pastagem ...") to its CLASS_ORDER name,
# so xlsx row labels (which may carry encoding/whitespace quirks) resolve to the
# canonical strings the rest of the stack uses.
_NUM_TO_CLASS = {str(i + 1): name for i, name in enumerate(CLASS_ORDER)}
YEARS = list(range(2008, 2025))


def _leading_num(label) -> str | None:
    if label is None:
        return None
    m = re.match(r"\s*(\d+)", str(label))
    return m.group(1) if m else None


def _find_workbook(src: Path, rgint_id: str) -> Path | None:
    if rgint_id in GOLDEN_IDS:
        golden = src / f"GOLDEN_ILUC_15_Classes_RGINT_{rgint_id}.xlsx"
        if golden.exists():
            return golden
    hits = sorted((src / "ALL_RGINTS").glob(f"*RGINT{rgint_id}_*.xlsx"))
    return hits[0] if hits else None


def _sheet_matrix(ws) -> dict[str, dict[str, float]]:
    """One sheet -> {origem_canonico: {destino_canonico: area_ha}} (skips zeros)."""
    rows = list(ws.iter_rows(values_only=True))
    matrix: dict[str, dict[str, float]] = {}
    for i in range(1, 16):
        origem = _NUM_TO_CLASS.get(_leading_num(rows[i][0]))
        if origem is None:
            raise ValueError(f"unmapped row label {rows[i][0]!r} on '{ws.title}'")
        for j in range(1, 16):
            val = rows[i][j]
            if val in (None, "", 0, 0.0):
                continue
            try:
                area = float(val)
            except (TypeError, ValueError):
                continue
            if area == 0.0:
                continue
            destino = CLASS_ORDER[j - 1]  # columns are positional (header unreliable)
            matrix.setdefault(origem, {})[destino] = area
    return matrix


def build_payload(wb_path: Path, meta: dict) -> dict:
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    matrices: dict[str, dict] = {}
    first_pair: dict[str, dict[str, float]] | None = None
    for sheet in wb.sheetnames:
        m = re.match(r"(\d{4})_(\d{4})", sheet)
        if not m:
            continue
        end_year = m.group(2)
        mat = _sheet_matrix(wb[sheet])
        matrices[end_year] = mat
        if m.group(1) == "2008":
            first_pair = mat
    wb.close()

    # Base-year 2008 stock as a diagonal = total area per origem in the 2008->2009
    # matrix (row sums). Falls back to empty if the 2008_2009 sheet is absent.
    diag: dict[str, dict[str, float]] = {}
    if first_pair:
        for origem, dests in first_pair.items():
            total = sum(dests.values())
            if total:
                diag[origem] = {origem: round(total, 4)}
    matrices["2008"] = diag

    return {
        "metadata": {
            "rgint": meta["rgint"],
            "nome": meta.get("nome", ""),
            "uf": meta.get("uf", ""),
            "biome": meta.get("biome", ""),
        },
        "anchor_years": [2008, 2017, 2024],
        "years": YEARS,
        "classes": list(CLASS_ORDER),
        "matrices": {str(y): matrices.get(str(y), {}) for y in YEARS},
    }


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("src", nargs="?",
                    default=os.environ.get("ILUC_MATRIX_SRC_DIR", str(DEFAULT_SRC)),
                    help="07_MATRIZES_15_CLASSES_FINAL directory")
    ap.add_argument("--out", default=str(RGINT_MATRIX_DIR),
                    help="output dir for per-region JSONs")
    args = ap.parse_args()

    src = Path(args.src)
    if not (src / "ALL_RGINTS").is_dir():
        sys.exit(f"source not found: {src}\\ALL_RGINTS")
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    index = json.loads(Path(INDEX_JSON).read_text(encoding="utf-8"))
    written, missing = 0, []
    for item in index:
        rid = str(item["rgint"])
        wb_path = _find_workbook(src, rid)
        if wb_path is None:
            missing.append(rid)
            continue
        payload = build_payload(wb_path, item)
        (out_dir / f"{rid}.json").write_text(
            json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
        src_kind = "GOLDEN" if wb_path.name.startswith("GOLDEN") else "raw"
        print(f"  {rid}: {wb_path.name}  ({src_kind})")
        written += 1

    print(f"\nWrote {written} matrices to {out_dir}")
    if missing:
        print(f"Missing source xlsx for {len(missing)} regions: {missing}")


if __name__ == "__main__":
    main()
